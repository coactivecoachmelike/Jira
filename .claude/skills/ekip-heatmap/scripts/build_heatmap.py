#!/usr/bin/env python3
"""Ürün Ekipleri Excel'inden board/ürün bazlı ekip heatmap'i üretir.

Kullanım:
  python build_heatmap.py --input "Ürün Ekipleri-2026-09-03.xlsx" \
      --output "heatmap-2026-09-03.xlsx" [--previous eski-heatmap.xlsx] [--aliases aliases.json]

--previous: Bir önceki heatmap (elle düzeltilmiş roller ve "Aktif mi?" değerleri buradan taşınır).
--aliases : {"Yanlış Yazım": "Doğru Yazım"} biçiminde isim düzeltme sözlüğü.
"""
import argparse
import difflib
import json
import os
import re
from collections import Counter, OrderedDict, defaultdict

import openpyxl
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SOURCE_SHEET = "Ürünler"
# Kaynak kolon -> rol adı. Sıra, bir satırdaki kişilerin heatmap'teki sırasını belirler.
ROLE_COLUMNS = OrderedDict([
    ("Scrum Master", "Scrum Master"),
    ("Product Manager", "Product Owner"),
    ("Yazılımcılar", "Yazılımcı"),
    ("Analistler", "Analist"),
    ("Testerlar", "Testçi"),
])
# Ana rol seçiminde eşitlik olursa öncelik sırası
ROLE_PRIORITY = ["Yazılımcı", "Analist", "Testçi", "Product Owner", "Scrum Master"]
EMPTY_TOKENS = {"", "-", ".", "?", "—", "nan", "none", "yok"}
UNKNOWN_ROLES = {None, "", "#N/A", "TBD"}

FILL_HEAD = PatternFill("solid", fgColor="CFE2F3")
FILL_KEY = PatternFill("solid", fgColor="B6D7A8")
FILL_TITLE = PatternFill("solid", fgColor="D9D9D9")
FILL_RED = PatternFill("solid", fgColor="F4C7C3")
FILL_YELLOW = PatternFill("solid", fgColor="FFFF00")
BOLD = Font(bold=True)

TR_LOWER = str.maketrans({"I": "ı", "İ": "i"})


def norm_key(name):
    return re.sub(r"\s+", " ", name.translate(TR_LOWER).lower()).strip()


def fold(name):
    table = str.maketrans("çğıöşüâî", "cgiosuai")
    return norm_key(name).translate(table)


def clean(v):
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip()


def split_people(cell):
    out = []
    if cell is None:
        return []
    for part in re.split(r"[,;\n]", str(cell)):
        p = clean(part).strip(" .")
        if p.lower() in EMPTY_TOKENS or not re.search(r"\w", p):
            continue
        out.append(p)
    return out


def read_source(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[SOURCE_SHEET]
    headers = [clean(c.value) for c in ws[1]]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if all(v in (None, "") for v in r):
            continue
        rows.append({h: v for h, v in zip(headers, r) if h})
    return wb, ws, headers, rows


def read_previous(path):
    """Önceki heatmap'ten isim->rol ve Board_Ürün->Aktif mi? haritalarını okur."""
    roles, active = {}, {}
    if not path:
        return roles, active
    wb = openpyxl.load_workbook(path, data_only=True)
    if "Final transposed_ürün heatmap" in wb.sheetnames:
        ws = wb["Final transposed_ürün heatmap"]
        for r in range(2, ws.max_row + 1):
            key = ws.cell(r, 6).value
            names = [ws.cell(r, c).value for c in range(7, ws.max_column + 1)]
            if not key or not any(isinstance(n, str) and n.strip() for n in names):
                continue
            if ws.cell(r, 5).value:
                active[key] = ws.cell(r, 5).value
            for c, n in enumerate(names, start=7):
                role = ws.cell(r + 1, c).value
                if isinstance(n, str) and n.strip() and isinstance(role, str) and role.strip() not in UNKNOWN_ROLES:
                    roles.setdefault(norm_key(n), role.strip())
    if "Takımlardaki kişiler unique" in wb.sheetnames:  # bu skill'in ürettiği dosya
        ws = wb["Takımlardaki kişiler unique"]
        if clean(ws.cell(1, 1).value) == "Kişi" and clean(ws.cell(1, 2).value) == "Rol":
            for r in range(2, ws.max_row + 1):
                n, role = ws.cell(r, 1).value, ws.cell(r, 2).value
                if n and role and role not in UNKNOWN_ROLES:
                    roles[norm_key(n)] = role
    return roles, active


def build(args):
    src_wb, src_ws, headers, rows = read_source(args.input)
    prev_roles, prev_active = read_previous(args.previous)
    aliases = {}
    alias_path = args.aliases or os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "aliases.json")
    if os.path.exists(alias_path):
        with open(alias_path, encoding="utf-8") as f:
            aliases = {norm_key(k): v for k, v in json.load(f).items() if not k.startswith("_")}

    issues = []
    used_aliases = Counter()
    spelling = defaultdict(Counter)  # norm_key -> yazım varyantları

    def canon(p):
        if norm_key(p) in aliases:
            used_aliases[(p, aliases[norm_key(p)])] += 1
            p = aliases[norm_key(p)]
        spelling[norm_key(p)][p] += 1
        return norm_key(p)

    products = []  # Board'u olan tüm satırlar
    for i, row in enumerate(rows, start=2):
        board = row.get("Board Name/Product Grouping")
        if board in (None, ""):
            continue
        urun = row.get("Ürün Adı") or row.get("Product") or ""
        key = f"{board}_{urun}"
        people = []  # (norm_key, rol)
        for col, role in ROLE_COLUMNS.items():
            raw = row.get(col)
            for p in split_people(raw):
                if len(p.split()) >= 4:
                    issues.append(("Ayraç eksik olabilir (tek hücrede 2 kişi?)", key, col, p))
                people.append((canon(p), role))
        products.append({"row": row, "key": key, "board": board, "people": people})

    display = {k: v.most_common(1)[0][0] for k, v in spelling.items()}
    for (old, new), n in used_aliases.items():
        issues.append(("İsim düzeltildi (aliases.json)", old, new, f"{n} hücre"))

    # Satır bazında unique kişiler (SM, PO, yazılımcı, analist, testçi sırasıyla)
    for p in products:
        seen = OrderedDict()
        for k, role in p["people"]:
            seen.setdefault(k, []).append(role)
        p["team"] = seen

    in_heatmap = [p for p in products if p["team"]]
    for p in products:
        if not p["team"]:
            issues.append(("Kişi yok, heatmap'e alınmadı", p["key"], "", ""))

    # Kişi bazında roller ve board'lar
    person_roles = defaultdict(Counter)
    person_keys = defaultdict(list)
    person_boards = defaultdict(set)
    for p in in_heatmap:
        for k, roles in p["team"].items():
            for r in roles:
                person_roles[k][r] += 1
            person_keys[k].append(p["key"])
            person_boards[k].add(clean(p["board"]))

    def main_role(k):
        if k in prev_roles:
            return prev_roles[k], "Önceki heatmap"
        c = person_roles[k]
        best = max(c.items(), key=lambda kv: (kv[1], -ROLE_PRIORITY.index(kv[0])))
        return best[0], "Türetildi"

    people = sorted(person_keys, key=lambda k: (-len(person_keys[k]), fold(display[k])))

    # Benzer isimler (olası yazım farkı)
    keys = sorted(person_keys, key=fold)
    for a_i, a in enumerate(keys):
        for b in keys[a_i + 1:]:
            if difflib.SequenceMatcher(None, fold(a), fold(b)).ratio() >= 0.87:
                issues.append(("Benzer isim (aynı kişi olabilir)", display[a], display[b], ""))

    # ------------------------------------------------------------------ çıktı
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # 1) Ürünler (kaynak kopyası)
    ws = wb.create_sheet("Ürünler")
    for r in src_ws.iter_rows(values_only=True):
        ws.append(list(r))
    for c in ws[1]:
        c.font = BOLD
        c.fill = FILL_TITLE
    ws.freeze_panes = "A2"

    # 2) Ürünler kopya
    ws = wb.create_sheet("Ürünler kopya")
    cols = ["Ekip", "Board Name/Product Grouping", "Type", "Product", "Ürün Adı", "Board_Ürün",
            "Board_Ürün_Value", "Own Solution", "Use Solution", "Scrum/Kanban", "Scrum Master",
            "Product Manager", "Yazılımcılar", "Analistler", "Testerlar", "Not", "TAKIM Formül",
            "TAKIM Value", "Heatmap'te mi?"]
    ws.append(cols)
    for i, p in enumerate(products, start=2):
        r = p["row"]
        vals = [r.get("Ekip"), r.get("Board Name/Product Grouping"), r.get("Type"), r.get("Product"),
                r.get("Ürün Adı") or r.get("Product"), f'=B{i}&"_"&E{i}', p["key"],
                r.get("Own Solution"), r.get("Use Solution"), r.get("Scrum/Kanban"),
                r.get("Scrum Master"), r.get("Product Manager"), r.get("Yazılımcılar"),
                r.get("Analistler"), r.get("Testerlar"), r.get("Not"),
                f'=K{i}&","&L{i}&","&M{i}&","&N{i}&","&O{i}',
                ",".join(clean(r.get(c)) for c in ROLE_COLUMNS),
                "Evet" if p["team"] else "Hayır (kişi yok)"]
        ws.append(vals)
    n_prod = len(products) + 1
    style_header(ws)

    # 3) Board list
    ws = wb.create_sheet("Board list")
    ws.append(["Board", "Ekip", "Ürün sayısı", "Unique kişi sayısı"])
    boards = OrderedDict()
    for p in products:
        boards.setdefault(p["board"], {"ekip": p["row"].get("Ekip"), "people": set()})
        boards[p["board"]]["people"].update(p["team"])
    for i, (b, d) in enumerate(boards.items(), start=2):
        ws.append([b, d["ekip"], f"=COUNTIF('Ürünler kopya'!$B$2:$B${n_prod},A{i})", len(d["people"])])
    style_header(ws)

    # 4) Board-ürün list
    ws = wb.create_sheet("Board-ürün list")
    max_team = max(len(p["team"]) for p in in_heatmap)
    last_col = get_column_letter(1 + max_team)
    cnt_col = 2 + max_team
    ws.append(["Board_Ürün_Value"] + ["TAKIM Value"] + [None] * (max_team - 1) + ["Kişi Sayısı"])
    for i, p in enumerate(in_heatmap, start=2):
        names = [display[k] for k in p["team"]]
        ws.cell(i, 1, p["key"])
        for j, n in enumerate(names, start=2):
            ws.cell(i, j, n)
        ws.cell(i, cnt_col, f"=COUNTA(B{i}:{last_col}{i})")
    n_bu = len(in_heatmap) + 1
    bu_range = f"'Board-ürün list'!$B$2:${last_col}${n_bu}"
    style_header(ws)

    # 5) Takımlardaki kişiler unique
    ws = wb.create_sheet("Takımlardaki kişiler unique")
    ws.append(["Kişi", "Rol", "Board_Ürün sayısı", "Farklı board sayısı", "Tüm roller",
               "Rol kaynağı", "Çalıştığı Board_Ürün'ler"])
    for i, k in enumerate(people, start=2):
        role, srcinfo = main_role(k)
        all_roles = ", ".join(f"{r} ({n})" for r, n in person_roles[k].most_common())
        ws.append([display[k], role, f"=COUNTIF({bu_range},A{i})", len(person_boards[k]),
                   all_roles, srcinfo, " | ".join(person_keys[k])])
    n_people = len(people) + 1
    style_header(ws)
    ws.conditional_formatting.add(f"C2:C{n_people}", heat_rule())
    uniq_range = f"'Takımlardaki kişiler unique'!$A$2:$C${n_people}"

    # 6) Final transposed_ürün heatmap (4 satırlık bloklar: isim / rol / board sayısı / boş)
    ws = wb.create_sheet("Final transposed_ürün heatmap")
    for c, h in enumerate(["PO", "SM", "PO - SM Aynı mı?", "Takım kişi sayısı", "Aktif mi?",
                           "Board_Ürün_Value"], start=1):
        ws.cell(1, c, h)
    kopya = f"'Ürünler kopya'!$G$2:$L${n_prod}"
    count_ranges, last = [], 7 + max_team - 1
    last_letter = get_column_letter(last)
    for b, p in enumerate(in_heatmap):
        r = 2 + 4 * b
        ws.cell(r, 1, f'=IFERROR(VLOOKUP(F{r},{kopya},6,FALSE)&"","")')
        ws.cell(r, 2, f'=IFERROR(VLOOKUP(F{r},{kopya},5,FALSE)&"","")')
        ws.cell(r, 3, f"=A{r}=B{r}")
        ws.cell(r, 4, f"=COUNTA(G{r}:{last_letter}{r})")
        ws.cell(r, 5, prev_active.get(p["key"]))
        ws.cell(r, 6, p["key"]).font = BOLD
        for j, k in enumerate(p["team"], start=7):
            L = get_column_letter(j)
            ws.cell(r, j, display[k])
            ws.cell(r + 1, j, f'=IFERROR(VLOOKUP({L}{r},{uniq_range},2,FALSE),"")')
            ws.cell(r + 2, j, f'=IFERROR(VLOOKUP({L}{r},{uniq_range},3,FALSE),"")')
        count_ranges.append(f"G{r + 2}:{get_column_letter(6 + len(p['team']))}{r + 2}")
        for rr in range(r, r + 4):
            for c in range(1, 6):
                ws.cell(rr, c).fill = FILL_HEAD
            ws.cell(rr, 6).fill = FILL_KEY
        ws.cell(r + 1, 6, "Rol").font = Font(italic=True, color="666666")
        ws.cell(r + 2, 6, "Kaç Board_Ürün'de?").font = Font(italic=True, color="666666")
    n_last = 1 + 4 * len(in_heatmap)
    for c in range(1, 7):
        ws.cell(1, c).font = BOLD
        ws.cell(1, c).fill = FILL_KEY if c == 6 else FILL_HEAD
    ws.conditional_formatting.add(" ".join(count_ranges), heat_rule())
    ws.conditional_formatting.add(f"D2:D{n_last}", CellIsRule(operator="greaterThanOrEqual",
                                                               formula=["9"], fill=FILL_RED))
    ws.conditional_formatting.add(f"C2:C{n_last}", FormulaRule(formula=["C2=TRUE"], fill=FILL_YELLOW))
    ws.freeze_panes = "G2"
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 19
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 11
    ws.column_dimensions["F"].width = 42
    for j in range(7, last + 1):
        ws.column_dimensions[get_column_letter(j)].width = 17
    for row in ws.iter_rows(min_row=1, max_row=n_last):
        for c in row:
            c.alignment = Alignment(vertical="center", wrap_text=c.column <= 6 and c.row == 1)

    # 7) Özet
    ws = wb.create_sheet("Özet", 0)
    ws.append(["Ekip Heatmap Özeti"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Kaynak dosya", args.input.split("/")[-1]])
    ws.append(["Heatmap'teki Board_Ürün sayısı", f"=COUNTA('Board-ürün list'!A2:A{n_bu})"])
    ws.append(["Farklı board sayısı", len({clean(p['board']) for p in in_heatmap})])
    ws.append(["Unique kişi sayısı", f"=COUNTA('Takımlardaki kişiler unique'!A2:A{n_people})"])
    ws.append(["5+ Board_Ürün'de çalışan kişi", f"=COUNTIF('Takımlardaki kişiler unique'!C2:C{n_people},\">=5\")"])
    ws.append([])
    ws.append(["En çok Board_Ürün'e dağılmış kişiler (darboğaz adayları)"])
    ws.cell(ws.max_row, 1).font = BOLD
    ws.append(["Kişi", "Rol", "Board_Ürün sayısı", "Farklı board sayısı"])
    hdr = ws.max_row
    for c in ws[hdr]:
        c.font = BOLD
        c.fill = FILL_HEAD
    for i in range(2, min(n_people, 21) + 1):
        ws.append([f"='Takımlardaki kişiler unique'!A{i}", f"='Takımlardaki kişiler unique'!B{i}",
                   f"='Takımlardaki kişiler unique'!C{i}", f"='Takımlardaki kişiler unique'!D{i}"])
    ws.conditional_formatting.add(f"C{hdr + 1}:C{ws.max_row}", heat_rule())
    ws.append([])
    ws.append(["Not: Board_Ürün sayısı, kişinin kaç ürün satırının ekibinde geçtiğini sayar; "
               "Farklı board sayısı aynı board'un ürünlerini tek sayar."])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 20

    # 8) Veri kalitesi
    ws = wb.create_sheet("Veri kalitesi")
    ws.append(["Kontrol", "Değer 1", "Değer 2", "Detay"])
    for it in issues:
        ws.append(list(it))
    style_header(ws)

    wb.save(args.output)
    print(json.dumps({
        "board_urun_in_heatmap": len(in_heatmap), "board_urun_total": len(products),
        "boards": len({clean(p['board']) for p in in_heatmap}), "people": len(people),
        "roles_from_previous": sum(1 for k in people if k in prev_roles),
        "issues": Counter(i[0] for i in issues),
        "top": [(display[k], main_role(k)[0], len(person_keys[k]), len(person_boards[k])) for k in people[:15]],
    }, ensure_ascii=False, indent=1))


def heat_rule():
    # 1 = yeşil, medyan = sarı, max = kırmızı (orijinal dosyadaki renkler)
    return ColorScaleRule(start_type="num", start_value=1, start_color="57BB8A",
                          mid_type="percentile", mid_value=50, mid_color="FFD666",
                          end_type="max", end_color="E67C73")


def style_header(ws):
    for c in ws[1]:
        c.font = BOLD
        c.fill = FILL_HEAD
    ws.freeze_panes = "B2"
    for col in ws.columns:
        L = col[0].column_letter
        width = max((len(str(c.value)) for c in col[:60] if c.value and not str(c.value).startswith("=")),
                    default=8)
        ws.column_dimensions[L].width = min(max(width + 2, 10), 45)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True)
    ap.add_argument("--output", required=True)
    ap.add_argument("--previous")
    ap.add_argument("--aliases")
    build(ap.parse_args())
